# выгрузка в Excel: листы с таблицами, на листе водопада MRR показатели считаются формулами

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import metrics

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="E8EEF7")


def _style_header(ws, row=1):
    for c in ws[row]:
        c.font = Font(name=FONT, bold=True, size=10)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")


def _set_font(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font.name != FONT or not c.font.bold:
                c.font = Font(name=FONT, size=10, bold=c.font.bold)


def _widths(ws, df=None, minw=10, maxw=48):
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 60)), start=1):
        w = max((len(str(c.value)) for c in col if c.value is not None), default=minw)
        ws.column_dimensions[get_column_letter(i)].width = max(minw, min(maxw, w + 2))


def _write_df(ws, df, formats=None, start_row=1):
    for j, name in enumerate(df.columns, start=1):
        ws.cell(row=start_row, column=j, value=str(name))
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, v in enumerate(row, start=1):
            if pd.isna(v):
                v = None
            elif hasattr(v, "item"):
                v = v.item()
            ws.cell(row=i, column=j, value=v)
    _style_header(ws, start_row)
    for col, fmt in (formats or {}).items():
        if col in df.columns:
            j = list(df.columns).index(col) + 1
            for i in range(start_row + 1, start_row + 1 + len(df)):
                ws.cell(row=i, column=j).number_format = fmt
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def waterfall_sheet(wb, w):
    ws = wb.create_sheet("MRR водопад")
    cols = ["month_end", "start_mrr", "new_mrr", "reactivation_mrr", "expansion_mrr", "price_uplift_mrr",
            "contraction_mrr", "churn_mrr", "end_mrr", "end_subscribers"]
    head = ["Конец месяца", "MRR на начало", "Новые", "Вернувшиеся", "Апгрейд", "Рост цены", "Даунгрейд", "Уход",
            "MRR на конец", "Подписчиков"]
    for j, h in enumerate(head, start=1):
        ws.cell(row=1, column=j, value=h)
    ws.cell(row=1, column=11, value="Проверка: начало плюс все компоненты минус конец")
    ws.cell(row=1, column=12, value="NRR")
    ws.cell(row=1, column=13, value="GRR")
    for i, row in enumerate(w[cols].itertuples(index=False), start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v.item() if hasattr(v, "item") else v)
        # контроль баланса и относительные метрики считаются формулами по этой же строке
        ws.cell(row=i, column=11, value=f"=ROUND(SUM(B{i}:H{i})-I{i},2)")
        ws.cell(row=i, column=12, value=f'=IF(B{i}=0,"",(B{i}+E{i}+F{i}+G{i}+H{i})/B{i})')
        ws.cell(row=i, column=13, value=f'=IF(B{i}=0,"",(B{i}+G{i}+H{i})/B{i})')
        for j in range(2, 10):
            ws.cell(row=i, column=j).number_format = "#,##0;-#,##0;-"
        ws.cell(row=i, column=11).number_format = "#,##0.00;-#,##0.00;-"
        ws.cell(row=i, column=12).number_format = "0.0%"
        ws.cell(row=i, column=13).number_format = "0.0%"
    last = len(w) + 1
    ws.cell(row=last + 2, column=1, value="Примечание: NRR не включает вернувшихся (реактивации), GRR не включает рост (апгрейд и цену).")
    _style_header(ws)
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "B2"
    return ws


def cohort_sheet(wb, coh):
    ws = wb.create_sheet("Когорты удержание")
    ret, rev, size = metrics.cohort_matrices(coh)
    ws.cell(row=1, column=1, value="Когорта (месяц первой оплаты)")
    ws.cell(row=1, column=2, value="Людей")
    for age in ret.columns:
        ws.cell(row=1, column=3 + int(age), value=f"мес {int(age)}")
    for i, idx in enumerate(ret.index, start=2):
        ws.cell(row=i, column=1, value=int(idx))
        ws.cell(row=i, column=2, value=int(size.loc[idx]))
        for age in ret.columns:
            v = ret.loc[idx, age]
            if pd.notna(v):
                c = ws.cell(row=i, column=3 + int(age), value=float(v) / 100)
                c.number_format = "0%"
    rng = f"C2:{get_column_letter(2 + ret.shape[1])}{1 + ret.shape[0]}"
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=0.3, start_color="FFFFFF",
                                                      end_type="num", end_value=1, end_color="2A78D6"))
    _style_header(ws)
    ws.freeze_panes = "C2"
    return ws


def build(res, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Описание"
    lines = [
        "Аналитика онлайн кинотеатра: выгрузка результатов",
        "Данные синтетические, их придумал генератор. Листы собраны автоматически скриптом run_all.py.",
        "",
        "Листы:",
        "KPI: главные цифры",
        "MRR водопад: движение MRR по месяцам, NRR и GRR посчитаны формулами",
        "Когорты удержание: доля людей, которые еще платят, по месяцам жизни",
        "Каналы: конверсия, цена регистрации, выручка за 180 дней",
        "Дожитие: медианы и доли живых подписок по каналам",
        "Модель риска: оценки логистической модели и значения, заложенные в генераторе",
        "Сегменты: портреты четырех сегментов зрителей",
        "Отток модель: качество модели и самые важные признаки",
        "Топ риска: 200 подписчиков с наибольшим шансом отмены в ближайшие 4 недели",
        "Эксперимент: итоги A/B теста и подгруппы",
        "Прогноз: регистрации на 28 дней вперед с интервалом",
        "Сверка с правдой: что оценили и что на самом деле заложено",
    ]
    for i, t in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=t)
    ws["A1"].font = Font(name=FONT, bold=True, size=13)
    ws.column_dimensions["A"].width = 100

    k = res["kpi"]
    ws = wb.create_sheet("KPI")
    rows = [("MRR на конец данных, руб", k["end_mrr"], "#,##0"), ("Платящих подписчиков", k["end_subscribers"], "#,##0"),
            ("Годовой NRR (произведение 12 месяцев)", k["nrr_12m"], "0.0%"),
            ("Медианный месячный GRR", k["median_monthly_grr"], "0.0%"),
            ("Удержание через 3 месяца, %", k["retention_m3"], "0.0"), ("Удержание через 6 месяцев, %", k["retention_m6"], "0.0"),
            ("Лучший канал по окупаемости", k["best_channel"], None), ("Худший канал по окупаемости", k["worst_channel"], None)]
    ws.append(["Показатель", "Значение"])
    for name, v, fmt in rows:
        ws.append([name, v])
        if fmt:
            ws.cell(row=ws.max_row, column=2).number_format = fmt
    _style_header(ws)

    m = res["metrics"]
    waterfall_sheet(wb, m["waterfall"])
    cohort_sheet(wb, m["cohorts"])

    ws = wb.create_sheet("Каналы")
    f = m["funnel"][["channel", "signups", "activated", "paying", "activation_pct", "conversion_pct", "spend",
                     "cac_per_signup", "revenue_180_per_signup", "revenue180_to_cac"]]
    f.columns = ["Канал", "Регистраций", "Активировались", "Платящих", "Активация, %", "Конверсия в платящих, %", "Затраты, руб",
                 "Цена регистрации, руб", "Выручка за 180 дней на регистрацию, руб", "Окупаемость"]
    _write_df(ws, f, {"Затраты, руб": "#,##0", "Регистраций": "#,##0", "Цена регистрации, руб": "#,##0",
                      "Выручка за 180 дней на регистрацию, руб": "#,##0", "Окупаемость": "0.0"})

    ws = wb.create_sheet("Дожитие")
    s = res["survival"]["km_summary"].copy()
    s.columns = ["Канал", "Подписок", "Медиана, недель", "Живы через 13 недель", "Через 26", "Через 52"]
    _write_df(ws, s, {"Живы через 13 недель": "0%", "Через 26": "0%", "Через 52": "0%"})

    ws = wb.create_sheet("Модель риска")
    h = res["survival"]["hazard_naive"].copy()
    truth = res["truth"]["hazard_ratios"]["planted_log_hazard"]
    h["заложено"] = h["term"].map(truth)
    h = h[["term", "coef", "se", "lo", "hi", "hazard_ratio", "заложено"]]
    h.columns = ["Признак", "Коэффициент (лог риска)", "Ошибка", "Нижняя граница", "Верхняя граница", "Отношение рисков", "Заложено"]
    _write_df(ws, h, {c: "0.000" for c in h.columns[1:]})

    ws = wb.create_sheet("Сегменты")
    p = res["segments"]["profile"].copy()
    p.insert(0, "Сегмент", [res["segments"]["names"][i] for i in p.index])
    p.columns = ["Сегмент", "Часов в неделю", "Часов за сессию", "Названий за сессию", "Доля выходных", "Разброс по неделям",
                 "Доля людей", "Людей"]
    _write_df(ws, p, {"Часов в неделю": "0.0", "Часов за сессию": "0.00", "Названий за сессию": "0.00",
                      "Доля выходных": "0%", "Разброс по неделям": "0.00", "Доля людей": "0%", "Людей": "#,##0"})

    ws = wb.create_sheet("Отток модель")
    r = res["churn"]["results"].copy()
    r.columns = ["Модель", "AUC", "PR AUC", "Brier", "LogLoss", "Доля уходящих на тесте"]
    _write_df(ws, r, {c: "0.000" for c in r.columns[1:]})
    imp = res["churn"]["importance"].head(15).copy()
    imp.columns = ["Признак", "Падение AUC при перемешивании", "Разброс"]
    start = len(r) + 4
    ws.cell(row=start - 1, column=1, value="Важность признаков (перестановочная)")
    ws.cell(row=start - 1, column=1).font = Font(name=FONT, bold=True)
    _write_df(ws, imp, {"Падение AUC при перемешивании": "0.0000", "Разброс": "0.0000"}, start_row=start)
    ws.freeze_panes = None

    ws = wb.create_sheet("Топ риска")
    t = res["top_risk"].copy()
    t.columns = ["user_id", "Риск отмены за 4 недели", "Часы: 2 недели к прошлым 8", "Часов за 2 недели", "Возраст подписки, недель",
                 "Тариф basic", "Тариф premium"]
    _write_df(ws, t, {"Риск отмены за 4 недели": "0.0%", "Часы: 2 недели к прошлым 8": "0.00", "Часов за 2 недели": "0.0",
                      "Возраст подписки, недель": "0"})

    ws = wb.create_sheet("Эксперимент")
    e = res["exp"]["itt"]
    ws.append(["Показатель", "Значение"])
    for name, v, fmt in [("Уходят без оффера", e["p_control"], "0.0%"), ("Уходят с оффером", e["p_offer"], "0.0%"),
                         ("Разница, п.п.", 100 * e["diff"], "0.0"), ("Интервал бутстреп, нижняя, п.п.", 100 * e["ci_boot"][0], "0.0"),
                         ("Интервал бутстреп, верхняя, п.п.", 100 * e["ci_boot"][1], "0.0"),
                         ("p перестановочного теста", e["perm_p"], "0.0000"), ("Людей в группе оффера", e["n_offer"], "#,##0"),
                         ("Людей в контроле", e["n_control"], "#,##0")]:
        ws.append([name, v])
        ws.cell(row=ws.max_row, column=2).number_format = fmt
    _style_header(ws)
    sub = res["exp"]["by_ratio"].copy()
    sub.columns = ["Падение часов (терциль)", "Людей", "Уходят без оффера", "Эффект", "Нижняя", "Верхняя"]
    start = ws.max_row + 3
    ws.cell(row=start - 1, column=1, value="Эффект по терцилям падения часов (доли, не п.п.)")
    ws.cell(row=start - 1, column=1).font = Font(name=FONT, bold=True)
    _write_df(ws, sub, {"Уходят без оффера": "0.0%", "Эффект": "0.0%", "Нижняя": "0.0%", "Верхняя": "0.0%"}, start_row=start)
    ws.freeze_panes = None

    ws = wb.create_sheet("Прогноз")
    fc = res["forecast"]["final"].copy()
    fc["date"] = fc["date"].dt.strftime("%Y-%m-%d")
    fc.columns = ["Дата", "Прогноз", "Нижняя граница 90%", "Верхняя граница 90%"]
    _write_df(ws, fc, {"Прогноз": "0.0", "Нижняя граница 90%": "0.0", "Верхняя граница 90%": "0.0"})
    n = len(fc) + 1
    ws.cell(row=n + 2, column=1, value="Итого за 28 дней")
    ws.cell(row=n + 2, column=2, value=f"=SUM(B2:B{n})").number_format = "#,##0"
    ws.cell(row=n + 3, column=1, value="Лучший метод по скользящей проверке")
    ws.cell(row=n + 3, column=2, value=res["forecast"]["best"])

    ws = wb.create_sheet("Сверка с правдой")
    sb = res["scoreboard"].copy()
    sb["covered"] = sb["covered"].map(lambda v: "" if v is None or pd.isna(v) else ("да" if v else "нет"))
    sb.columns = ["Блок", "Что оценивали", "Оценка", "Нижняя", "Верхняя", "Заложено", "Единицы", "Правда внутри интервала"]
    _write_df(ws, sb, {c: "0.000" for c in ["Оценка", "Нижняя", "Верхняя", "Заложено"]})

    for ws in wb.worksheets:
        _widths(ws)
        _set_font(ws)
    wb["Описание"].column_dimensions["A"].width = 100
    wb["Описание"]["A1"].font = Font(name=FONT, bold=True, size=13)
    wb.calculation.fullCalcOnLoad = True      # Excel пересчитает формулы при открытии
    wb.save(path)
    return path
